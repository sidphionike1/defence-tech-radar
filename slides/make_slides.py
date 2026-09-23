r"""Generate the 2-slide showcase deck (slides/radar-demo-slides.pptx).

Rerun any time:  ..\backend\.venv\Scripts\python.exe make_slides.py
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Radar-console palette (matches the frontend theme).
BG = RGBColor(0x03, 0x0A, 0x06)
PANEL = RGBColor(0x07, 0x19, 0x0F)
GREEN = RGBColor(0x43, 0xFF, 0x9A)
WHITE = RGBColor(0xEA, 0xFF, 0xF4)
DIM = RGBColor(0x8D, 0x9B, 0x95)
RED = RGBColor(0xFF, 0x55, 0x77)
BLUE = RGBColor(0x4D, 0xAB, 0xF7)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)  # rectangle
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def text_box(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def para(tf, text, size, color=WHITE, bold=False, first=False, align=PP_ALIGN.LEFT,
         font="Consolas", space_after=6, bullet_color=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if bullet_color is not None:
        run = p.add_run()
        run.text = "▸ "
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = bullet_color
        run.font.name = font
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


def panel(slide, left, top, width, height, edge=GREEN):
    shp = slide.shapes.add_shape(1, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL
    shp.line.color.rgb = edge
    shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


# ---------------------------------------------------------------- slide 1
s = add_slide()

tf = text_box(s, Inches(0.6), Inches(0.45), Inches(12.1), Inches(1.6))
para(tf, "LO-TARGET RADAR", 40, GREEN, bold=True, first=True, space_after=2)
para(tf, "Low-observable detection — multistatic sensor fusion + ML + Claude sitreps", 17, DIM)

tf = text_box(s, Inches(0.6), Inches(1.9), Inches(12.1), Inches(1.0))
para(tf, "Conventional radar applies ONE fixed threshold to ONE sensor.", 21, WHITE, bold=True, first=True, space_after=2)
para(tf, "Small drones and stealth-shaped targets sit at the noise floor — and simply never appear.", 17, DIM)

# Left panel — the miss
p1 = panel(s, Inches(0.6), Inches(3.15), Inches(5.9), Inches(2.9), edge=RED)
tf = p1.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.28)
tf.margin_top = Inches(0.2)
para(tf, "CONVENTIONAL THRESHOLD", 15, RED, bold=True, first=True, space_after=8)
para(tf, "detected = primary SNR > 10 dB", 14, WHITE, space_after=8)
para(tf, "Stealth contact at 2–8 dB → below threshold", 13, DIM, bullet_color=RED)
para(tf, "0 of 6 low-observable targets detected", 13, DIM, bullet_color=RED)
para(tf, "A permanent blind spot — by design", 13, DIM, bullet_color=RED)

# Right panel — the catch
p2 = panel(s, Inches(6.85), Inches(3.15), Inches(5.9), Inches(2.9), edge=GREEN)
tf = p2.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.28)
tf.margin_top = Inches(0.2)
para(tf, "FUSION + ML (this project)", 15, GREEN, bold=True, first=True, space_after=8)
para(tf, "3 separated receivers · RCS is angle-dependent", 13, WHITE, bullet_color=GREEN)
para(tf, "Fused confidence — same 10 dB-equivalent rule", 13, WHITE, bullet_color=GREEN)
para(tf, "RandomForest classifies · IsolationForest flags anomalies", 13, WHITE, bullet_color=GREEN)
para(tf, "Claude writes the operator's situation report on click", 13, WHITE, bullet_color=GREEN)

tf = text_box(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.8))
para(tf, "“No single sensor was confident. Combined, they are.”", 20, GREEN, bold=True,
     first=True, align=PP_ALIGN.CENTER)
para(tf, "The same math behind real counter-stealth systems — and fraud, intrusion & IoT anomaly detection.",
     13, DIM, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- slide 2
s = add_slide()

tf = text_box(s, Inches(0.6), Inches(0.45), Inches(12.1), Inches(1.2))
para(tf, "LIVE: SINGLE THRESHOLD MISSES IT — FUSION + ML CATCHES IT", 26, GREEN, bold=True, first=True, space_after=2)
para(tf, "46 seeded targets · two scopes, same airspace · end-to-end pipeline built in under an hour", 15, DIM)

# Stat tiles
stats = [
    ("0%", "stealth targets caught by\nconventional threshold", RED),
    ("100%", "caught by sensor fusion\n(6 of 6 recovered)", GREEN),
    ("<3 s", "click-to-explain — Claude sitrep,\nfallback if offline", BLUE),
]
x = Inches(0.6)
for big, small, color in stats:
    tile = panel(s, x, Inches(1.75), Inches(3.94), Inches(1.85), edge=color)
    tf = tile.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.12)
    para(tf, big, 44, color, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=2)
    for line in small.split("\n"):
        para(tf, line, 12.5, WHITE, align=PP_ALIGN.CENTER, space_after=0)
    x += Inches(4.11)

# How it's built
p3 = panel(s, Inches(0.6), Inches(3.95), Inches(5.9), Inches(2.6))
tf = p3.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.28)
tf.margin_top = Inches(0.18)
para(tf, "HOW IT WORKS", 14, GREEN, bold=True, first=True, space_after=6)
para(tf, "numpy generator (seed 42): 40 normal + 6 stealth targets", 12.5, WHITE, bullet_color=GREEN)
para(tf, "FastAPI serves baseline vs fused detections, polled at 1.5 s", 12.5, WHITE, bullet_color=GREEN)
para(tf, "RandomForest + IsolationForest train at startup in <2 s", 12.5, WHITE, bullet_color=GREEN)
para(tf, "Startup assertion guarantees the recovery moment exists", 12.5, WHITE, bullet_color=GREEN)
para(tf, "Claude (claude-sonnet-4-6) explains any contact on click", 12.5, WHITE, bullet_color=GREEN)

# Parallel build / mock story
p4 = panel(s, Inches(6.85), Inches(3.95), Inches(5.9), Inches(2.6), edge=BLUE)
tf = p4.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.28)
tf.margin_top = Inches(0.18)
para(tf, "BUILT IN PARALLEL, INTEGRATED IN MINUTES", 14, BLUE, bold=True, first=True, space_after=6)
para(tf, "API contract locked in the first 5 minutes", 12.5, WHITE, bullet_color=BLUE)
para(tf, "Frontend built on a seeded in-browser mock of the backend", 12.5, WHITE, bullet_color=BLUE)
para(tf, "Real captured backend output kept in-repo as ground truth", 12.5, WHITE, bullet_color=BLUE)
para(tf, "Integration = one flag flip (USE_MOCK) — not a rebuild", 12.5, WHITE, bullet_color=BLUE)
para(tf, "?api=mock still runs the full UI with zero backend", 12.5, WHITE, bullet_color=BLUE)

tf = text_box(s, Inches(0.6), Inches(6.7), Inches(12.1), Inches(0.6))
para(tf, "Pranav — backend/ML · Sid — frontend · Shubhangi — docs, QA & demo   |   FastAPI · scikit-learn · Canvas/vanilla JS · Claude API",
     13, DIM, first=True, align=PP_ALIGN.CENTER)

out = "radar-demo-slides.pptx"
prs.save(out)
print("wrote", out)
